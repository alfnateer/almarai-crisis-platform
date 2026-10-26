#!/usr/bin/env python3
"""Regenerate the website's template data from the BCM Toolkit workbook.

Run after replacing template/BCM_Implementation_Toolkit_*.xlsx:

    python3 tools/build_template_data.py template/BCM_Implementation_Toolkit_v6.4_Generic.xlsx

Writes:
  js/template.js       - the workbook itself (base64), used as the export template
  js/template-data.js  - static labels / defaults, ISO-map status wiring and the
                         read-only reference pages (Instructions, ERM RA Guidance)
"""
import base64
import datetime as dt
import html
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parent.parent


def plain(v):
    if v is None:
        return None
    if hasattr(v, "text"):  # array formula
        return None
    if isinstance(v, str) and v.startswith("="):
        return None
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.date):
        return v.isoformat()
    return v


def formula(v):
    if hasattr(v, "text"):
        return v.text
    if isinstance(v, str) and v.startswith("="):
        return v
    return None


def static_values(wb):
    out = {}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                p = plain(c.value)
                if p is not None and p != "":
                    cells[c.coordinate] = p
        out[ws.title] = cells
    return out


def formula_cells(wb):
    """Addresses of formula cells per sheet (the website never writes these)."""
    out = {}
    for ws in wb.worksheets:
        out[ws.title] = sorted(
            c.coordinate for row in ws.iter_rows() for c in row if formula(c.value)
        )
    return out


def iso_map_wiring(wb):
    ws = wb["19. ISO 22301 Map"]
    rows = {}
    for r in range(5, 44):
        f = formula(ws[f"D{r}"].value)
        if not f:
            continue
        ext = re.search(r'"(EXT-\d+)"', f)
        pdca = re.search(r"'17\. BCMS Progress \(PDCA\)'!\$F(\d+)", f)
        if "'00. Document Control'" in f:
            rows[r] = {"type": "doc75"}
        elif "_xlpm.roles" in f:
            rows[r] = {"type": "roles53"}
        elif "$B$18:$D$25" in f:
            rows[r] = {"type": "roles51"}
        elif "_xlpm.p," in f and ext and pdca:
            rows[r] = {"type": "pdcaext", "row": int(pdca.group(1)), "ext": ext.group(1)}
        elif ext:
            rows[r] = {"type": "ext", "ext": ext.group(1)}
        elif pdca:
            rows[r] = {"type": "pdca", "row": int(pdca.group(1))}
        else:
            raise SystemExit(f"Unrecognised ISO map status formula in D{r}: {f[:80]}")
    i_formula = [r for r in range(5, 44) if formula(ws[f"I{r}"].value)]
    return {"D": rows, "Iformula": i_formula}


def colour(c):
    try:
        if c.fill and c.fill.fill_type == "solid":
            rgb = c.fill.fgColor.rgb
            if isinstance(rgb, str) and len(rgb) == 8 and rgb != "00000000":
                return "#" + rgb[2:]
    except Exception:
        pass
    return None


def font_colour(c):
    try:
        rgb = c.font.color.rgb if c.font and c.font.color else None
        if isinstance(rgb, str) and len(rgb) == 8:
            return "#" + rgb[2:]
    except Exception:
        pass
    return None


def render_sheet_html(ws, max_col=None):
    """Render a read-only sheet as an HTML table (merges, fills, bold)."""
    max_row = ws.max_row
    max_col = max_col or ws.max_column
    covered, spans = set(), {}
    for m in ws.merged_cells.ranges:
        spans[(m.min_row, m.min_col)] = (m.max_row - m.min_row + 1, m.max_col - m.min_col + 1)
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if (r, c) != (m.min_row, m.min_col):
                    covered.add((r, c))
    used_cols = [
        c for c in range(1, max_col + 1)
        if any(ws.cell(r, c).value is not None for r in range(1, max_row + 1))
    ]
    first_col = min(used_cols) if used_cols else 1
    last_col = max(used_cols) if used_cols else 1
    widths = []
    for c in range(first_col, last_col + 1):
        w = ws.column_dimensions[get_column_letter(c)].width or 10
        widths.append(max(4, min(80, w)))
    total = sum(widths)
    out = ['<table class="ref-table"><colgroup>']
    for w in widths:
        out.append(f'<col style="width:{w / total * 100:.2f}%">')
    out.append("</colgroup><tbody>")
    for r in range(1, max_row + 1):
        if all(ws.cell(r, c).value is None for c in range(first_col, last_col + 1)) and not any(
            (r, c) in spans for c in range(first_col, last_col + 1)
        ):
            continue
        out.append("<tr>")
        for c in range(first_col, last_col + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(r, c)
            attrs, style = [], []
            if (r, c) in spans:
                rs, cs = spans[(r, c)]
                # Clamp spans to the rendered area.
                cs = min(cs, last_col - c + 1)
                if rs > 1:
                    attrs.append(f'rowspan="{rs}"')
                if cs > 1:
                    attrs.append(f'colspan="{cs}"')
            bg = colour(cell)
            if bg:
                style.append(f"background:{bg}")
                fc = font_colour(cell)
                if fc and fc.lower() != "#000000":
                    style.append(f"color:{fc}")
            if cell.font and cell.font.b:
                style.append("font-weight:600")
            v = cell.value
            if formula(v):
                text = ""
            elif isinstance(v, (dt.date, dt.datetime)):
                text = v.strftime("%Y-%m-%d")
            elif v is None:
                text = ""
            else:
                text = str(v)
            if style:
                attrs.append('style="' + ";".join(style) + '"')
            body = html.escape(text).replace("\n", "<br>")
            out.append(f"<td {' '.join(attrs)}>{body}</td>")
        out.append("</tr>")
    out.append("</tbody></table>")
    return "".join(out)


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else next((ROOT / "template").glob("*.xlsx"))
    wb = openpyxl.load_workbook(src)
    data = {
        "source": src.name,
        "sheets": [ws.title for ws in wb.worksheets],
        "static": static_values(wb),
        "formulas": formula_cells(wb),
        "isoMap": iso_map_wiring(wb),
        "reference": {
            "01. Instructions": render_sheet_html(wb["01. Instructions"]),
            "20. ERM RA Guidance": render_sheet_html(wb["20. ERM RA Guidance"]),
        },
    }
    (ROOT / "js" / "template-data.js").write_text(
        "/* Generated by tools/build_template_data.py - do not edit by hand. */\n"
        "window.BCM_TEMPLATE_DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    b64 = base64.b64encode(src.read_bytes()).decode()
    (ROOT / "js" / "template.js").write_text(
        "/* Generated by tools/build_template_data.py - the original workbook, used as the export template. */\n"
        f"window.BCM_TEMPLATE_NAME = {json.dumps(src.name)};\n"
        f'window.BCM_TEMPLATE_B64 = "{b64}";\n',
        encoding="utf-8",
    )
    print("wrote js/template-data.js and js/template.js from", src.name)


if __name__ == "__main__":
    main()
